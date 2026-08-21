#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EN_writer.py — Ghi từ mới xuống CUỐI sheet từ vựng, giữ nguyên dữ liệu cũ
=========================================================================
Dùng chung cho EN_video.py và EN_add.py.
"""

import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from EN_common import (
    C_AUDIO, C_CHECK, C_DATE, C_DAY1, C_EX, C_EXVI, C_IPA, C_LEVEL, C_NOTE,
    C_ROOT, C_STT, C_TOPIC, C_TRY1, C_TRY2, C_VIET, C_WORD, C_YOUG,
    FIRST_DATA_ROW, HEADER_ROW, HISTORY_START_ROW, N_COLS, ROW_HEIGHT,
    SHEET_VOCAB, audio_url, dict_url, log, youglish_url,
)


def safe_save(wb, path, quiet=False):
    """Lưu file; nếu Excel đang mở file thì báo rõ cho người dùng."""
    try:
        wb.save(path)
        return True
    except PermissionError:
        if not quiet:
            log("")
            log("  [X] KHÔNG LƯU ĐƯỢC — file đang mở trong Excel.")
            log("      ➜ Đóng Excel rồi chạy lại.")
        return False


def last_data_row(ws):
    last = HEADER_ROW
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(r, C_WORD).value:
            last = r
    return last


def existing_words(ws, last):
    """Tập hợp từ đã có (chữ thường) để loại trùng."""
    out = set()
    for r in range(FIRST_DATA_ROW, last + 1):
        v = ws.cell(r, C_WORD).value
        if v:
            out.add(str(v).strip().lower())
    return out


def clone_style(ws, src_row, dst_row):
    for c in range(1, N_COLS + 1):
        ws.cell(dst_row, c)._style = ws.cell(src_row, c)._style
    ws.row_dimensions[dst_row].height = (
        ws.row_dimensions[src_row].height or ROW_HEIGHT)


def bump_ranges(ws, old_last, new_last):
    """Nới rộng công thức thống kê hàng 1, auto-filter và định dạng có điều kiện."""
    for c in range(1, N_COLS + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str) and v.startswith("="):
            def rep(m):
                n = int(m.group(3))
                return (f"{m.group(1)}:{m.group(2)}{max(n, new_last)}"
                        if n >= old_last - 5 else m.group(0))
            ws.cell(1, c).value = re.sub(
                r"(\$?[A-Z]{1,2}\$?\d+):(\$?[A-Z]{1,2}\$?)(\d+)", rep, v)
    try:
        ws.auto_filter.ref = (f"A{HEADER_ROW}:"
                              f"{get_column_letter(N_COLS)}{new_last}")
    except Exception:
        pass
    try:
        cf = ws.conditional_formatting
        for rng in list(cf._cf_rules.keys()):
            if str(rng.sqref).startswith("H"):
                rules = cf._cf_rules.pop(rng)
                rng.sqref = openpyxl.worksheet.cell_range.MultiCellRange(
                    f"H{FIRST_DATA_ROW}:H{new_last}")
                cf._cf_rules[rng] = rules
    except Exception:
        pass


def write_words(ws, words, tpl_row, level="VIDEO",
                topic_title="", topic_url="", source=""):
    """
    words: list dict {word, ipa, viet, root, topic, ex, ex_vi, freq, note}
    Ghi từ dòng tpl_row+1 trở đi. Trả (start_row, end_row).
    """
    start = tpl_row + 1
    max_stt = 0
    for r in range(FIRST_DATA_ROW, tpl_row + 1):
        v = ws.cell(r, C_STT).value
        if isinstance(v, (int, float)):
            max_stt = max(max_stt, int(v))
    today = f"{datetime.now():%Y-%m-%d}"

    for i, w in enumerate(words):
        r = start + i
        clone_style(ws, tpl_row, r)

        ws.cell(r, C_STT, max_stt + 1 + i)
        ws.cell(r, C_LEVEL, w.get("level") or level)
        ws.cell(r, C_WORD, w["word"])
        ip = ws.cell(r, C_IPA, w.get("ipa") or "🔗 tra IPA")
        if not w.get("ipa"):
            ip.hyperlink = dict_url(w["word"])
            ip.font = Font(size=10, color="1565C0", underline="single")
        ws.cell(r, C_VIET, w.get("viet") or "")
        ws.cell(r, C_TRY1, None)
        ws.cell(r, C_TRY2, None)
        ws.cell(r, C_CHECK,
                f'=IF(TRIM($G{r})="","",IF(EXACT(LOWER(TRIM($G{r})),'
                f'LOWER(TRIM($C{r}))),"✅ Đúng","❌ Sai"))')

        a = ws.cell(r, C_AUDIO, "▶")
        a.hyperlink = audio_url(w["word"])
        a.font = Font(size=13, color="1565C0")
        y = ws.cell(r, C_YOUG, "🔍")
        y.hyperlink = youglish_url(w["word"])
        y.font = Font(size=13, color="1565C0")

        for c in range(C_DAY1, C_DAY1 + 7):
            ws.cell(r, c, None)

        ws.cell(r, C_DATE, today)

        note = w.get("note") or ""
        if not note and w.get("freq"):
            note = f"📹 Từ video · xuất hiện {w['freq']} lần"
            if source:
                note += f" · nguồn: {source}"
        ws.cell(r, C_NOTE, note)

        ws.cell(r, C_ROOT, w.get("root") or "")

        t = ws.cell(r, C_TOPIC, w.get("topic") or (f"🎬 {topic_title}"
                                                   if topic_title else ""))
        if topic_url and not w.get("topic"):
            t.hyperlink = topic_url
            t.font = Font(size=9, bold=True, color="1565C0",
                          underline="single")

        ws.cell(r, C_EX, w.get("ex") or "")
        ws.cell(r, C_EXVI, w.get("ex_vi") or "")

    return start, start + len(words) - 1


def append_history(ws, title, url, n_new, n_dup):
    r = HISTORY_START_ROW
    while ws.cell(r, 1).value:
        r += 1
    ws.cell(r, 1, f"{datetime.now():%Y-%m-%d %H:%M}")
    ws.cell(r, 2, title)
    c = ws.cell(r, 3, "▶ Mở video")
    if url:
        c.hyperlink = url
    c.font = Font(color="1565C0", underline="single")
    ws.cell(r, 4, n_new).alignment = Alignment(horizontal="center")
    ws.cell(r, 5, n_dup).alignment = Alignment(horizontal="center")


def get_vocab_sheet(wb):
    if SHEET_VOCAB in wb.sheetnames:
        return wb[SHEET_VOCAB]
    for name in wb.sheetnames:
        if "Từ Vựng" in name or "Tu Vung" in name:
            return wb[name]
    raise SystemExit(f"[X] Không tìm thấy sheet '{SHEET_VOCAB}' trong file.")
